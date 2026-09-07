#!/usr/bin/env python3
"""
Verwerkt de datering-data van de domeinexpert (data/Begraafplaatsen
Zuid-Holland - datering.xlsx, tabblad "Alle begraafplaatsen", 448 rijen)
tot een aparte referentielaag voor de viewer.

Waarom een aparte laag i.p.v. een veld op de bestaande 448 begraafplaatsen:
geprobeerd is eerst een 1-op-1 koppeling te maken (naam+plaats, daarna
adres-geocodering + ruimtelijke matching tegen de bestaande ingang-punten).
De ruimtelijke matching werkt zelf goed (90% binnen 150m), maar legt bloot
dat de 448 xlsx-rijen NIET 1-op-1 overeenkomen met onze 448 terreinen: een
deel is een sub-onderdeel van een terrein dat wij als één geheel tellen (bv.
een Joodse afdeling binnen een algemene begraafplaats, letterlijk dezelfde
coördinaat), en een ander deel is een begraafplaats die niet in onze 448
zit. Een 1-op-1 koppeling zou dus op ruim 20 plekken een verkeerde of
misleidende suggestie van precisie geven -- zelfde valkuil als de
in_hoofddataset-heuristiek die eerder al bij de verdwenen-begraafplaatsen-
laag is verwijderd (2026-08-31, zie docs/geschiedenis.md). Deze laag toont
daarom de datering-punten op zichzelf, gegeocodeerd op het eigen adres uit
de xlsx, zonder gepretendeerde koppeling aan de hoofddataset.

Geocoding via de gratis PDOK Locatieserver (Bezoekadres + Huisnummer + PC +
Plaats -> centroide_ll van het beste treffer, meestal postcode-niveau,
ruim voldoende nauwkeurig om het juiste terrein te tonen).

De bron-xlsx zelf staat NIET in de repository: naast de datering-kolommen
bevat het werkblad ook Eigenaar/Contactpersoon/Telefoon/E-mail per
begraafplaats (persoonsgegevens van derden, niet van dit project). Alleen
het resultaat van dit script (data/generated/datering.geojson, dat enkel
naam/plaats/gemeente/jaartal/periode bevat) wordt gecommit. Om dit script
opnieuw te draaien is de originele xlsx van de domeinexpert lokaal nodig op
het pad hieronder (XLSX_PATH).

Output:
  data/generated/datering.geojson
"""
from __future__ import annotations

import json
import re
import time
import urllib.parse
import urllib.request
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = REPO_ROOT / "data" / "Begraafplaatsen Zuid-Holland - datering.xlsx"
GENERATED_DIR = REPO_ROOT / "data" / "generated"
OUT_PATH = GENERATED_DIR / "datering.geojson"

GEOCODE_URL = "https://api.pdok.nl/bzk/locatieserver/search/v3_1/free"
POINT_PATTERN = re.compile(r"POINT\(([-\d.]+) ([-\d.]+)\)")

# Kolomnummers (1-based) in het tabblad "Alle begraafplaatsen" -- de kop
# "Plaats" komt twee keer voor (bezoekadres en postadres van de eigenaar),
# dus positioneel gebruiken i.p.v. op kolomnaam opzoeken.
COL_NAAM = 4
COL_BEZOEKADRES = 7
COL_HUISNUMMER = 8
COL_PC = 10
COL_PLAATS = 11
COL_GEMEENTE = 12
COL_JAARTAL = 26
COL_CIRCA = 27
COL_PERIODE = 28


def geocode(adres: str | None, huisnummer, postcode: str | None, plaats: str | None) -> tuple[float, float] | None:
    q_parts = [adres or "", str(huisnummer or ""), postcode or "", plaats or ""]
    q = " ".join(p for p in q_parts if p).strip()
    if not q:
        return None
    url = f"{GEOCODE_URL}?{urllib.parse.urlencode({'q': q, 'rows': 1})}"
    with urllib.request.urlopen(url, timeout=10) as resp:
        data = json.loads(resp.read())
    docs = data.get("response", {}).get("docs", [])
    if not docs:
        return None
    m = POINT_PATTERN.match(docs[0].get("centroide_ll", ""))
    if not m:
        return None
    return float(m.group(1)), float(m.group(2))


def main() -> None:
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb["Alle begraafplaatsen"]

    punten = []
    geocode_mislukt = []
    for r in range(2, ws.max_row + 1):
        naam = ws.cell(row=r, column=COL_NAAM).value
        if not naam:
            continue
        adres = ws.cell(row=r, column=COL_BEZOEKADRES).value
        huisnummer = ws.cell(row=r, column=COL_HUISNUMMER).value
        pc = ws.cell(row=r, column=COL_PC).value
        plaats = ws.cell(row=r, column=COL_PLAATS).value
        gemeente = ws.cell(row=r, column=COL_GEMEENTE).value
        jaartal = ws.cell(row=r, column=COL_JAARTAL).value
        circa = ws.cell(row=r, column=COL_CIRCA).value == "Ja"
        periode = ws.cell(row=r, column=COL_PERIODE).value

        coords = geocode(adres, huisnummer, pc, plaats)
        time.sleep(0.03)  # de PDOK Locatieserver is een gratis, gedeelde dienst
        if coords is None:
            geocode_mislukt.append(f"{naam} ({plaats})")
            continue
        lon, lat = coords

        punten.append({
            "naam": naam,
            "plaats": plaats,
            "gemeente": gemeente,
            "jaartal": jaartal,
            "jaartal_circa": circa,
            "periode": periode,
            "lon": round(lon, 6),
            "lat": round(lat, 6),
        })

    punten.sort(key=lambda k: (k["plaats"] or "", k["naam"]))

    n_jaartal = sum(1 for p in punten if p["jaartal"])
    n_periode = sum(1 for p in punten if p["periode"] and not p["jaartal"])

    out = {
        "type": "FeatureCollection",
        "waarschuwing": (
            "Bron: de domeinexpert (data/Begraafplaatsen Zuid-Holland - "
            "datering.xlsx). Zelfstandige laag, geen 1-op-1 koppeling aan "
            "de 448 begraafplaatsen van de hoofddataset -- zie de "
            "toelichting in scripts/build_datering.py en methode.html. "
            "Locatie is gegeocodeerd op het bezoekadres (PDOK Locatieserver), "
            "meestal postcode-nauwkeurig."
        ),
        "aantal": len(punten),
        "aantal_met_jaartal": n_jaartal,
        "aantal_met_alleen_periode": n_periode,
        "features": [
            {
                "type": "Feature",
                "properties": {k: v for k, v in punt.items() if k not in ("lon", "lat")},
                "geometry": {"type": "Point", "coordinates": [punt["lon"], punt["lat"]]},
            }
            for punt in punten
        ],
    }
    GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    OUT_PATH.write_text(json.dumps(out, indent=2, ensure_ascii=False), encoding="utf-8")

    print(f"datering geschreven -> {OUT_PATH}")
    print(f"  {len(punten)} punten, {n_jaartal} met jaartal, {n_periode} met alleen periode")
    if geocode_mislukt:
        print(f"  {len(geocode_mislukt)} niet gegeocodeerd: {geocode_mislukt}")


if __name__ == "__main__":
    main()
